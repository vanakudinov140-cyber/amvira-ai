from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)


def h2(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 2")


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def nums(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title(doc, "Финальное техническое задание")
    add_subtitle(doc, "Проект: VK-бот для розыгрыша в VR-клубе Avatar Arena Omsk VR")
    add_subtitle(doc, "Площадка: https://vk.ru/avatararenaomskvr")
    add_subtitle(doc, f"Дата: {date.today().strftime('%d.%m.%Y')}")
    doc.add_paragraph("")

    h2(doc, "1. Цель проекта")
    doc.add_paragraph(
        "Запустить понятную и прозрачную механику розыгрыша в сообществе VK, "
        "которая мотивирует гостей VR-клуба на повторные покупки и увеличивает активность аудитории."
    )

    h2(doc, "2. Ключевая механика (утвержденный формат)")
    bullets(
        doc,
        [
            "Розыгрыш проходит в 3 одинаковых этапа по 30 дней.",
            "1 QR-код = 1 шанс на выигрыш.",
            "Клиент получает несколько QR в зависимости от суммы чека: сумма / 1200 (целая часть).",
            "Пример: чек 12 000 руб. = 10 QR = 10 шансов.",
            "QR одноразовые: повторная активация невозможна.",
        ],
    )

    h2(doc, "3. Логика этапов")
    bullets(
        doc,
        [
            "В конце каждого этапа выбираются победители случайным образом из всех активных шансов.",
            "Коды, которые выиграли в этапе, исключаются из следующих этапов.",
            "Коды, которые не выиграли, автоматически участвуют в следующем этапе.",
            "Таким образом, участники сохраняют шанс на победу в следующих этапах.",
        ],
    )

    h2(doc, "4. Как работает сценарий для гостя")
    nums(
        doc,
        [
            "Гость получает QR-код(ы) после покупки.",
            "Сканирует QR и открывает диалог с VK-ботом.",
            "Бот автоматически фиксирует активацию и начисляет шанс.",
            "Гость получает подтверждение и видит, сколько шансов уже начислено.",
        ],
    )

    h2(doc, "5. Как работает розыгрыш")
    bullets(
        doc,
        [
            "Розыгрыш запускается внутри бота администратором, без сторонних сервисов.",
            "Бот использует внутреннюю базу активированных QR и шансов.",
            "Результаты фиксируются в истории: этап, дата, победители.",
            "Доступны выгрузки в CSV/Excel.",
        ],
    )

    h2(doc, "6. Административные функции")
    bullets(
        doc,
        [
            "Создание/закрытие этапов.",
            "Запуск выбора 1 или нескольких победителей.",
            "Просмотр списка участников и количества шансов.",
            "Выгрузка данных и истории розыгрышей.",
            "Повторный розыгрыш при необходимости.",
        ],
    )

    h2(doc, "7. Что входит в работу исполнителя")
    bullets(
        doc,
        [
            "Разработка VK-бота на Python под утвержденную механику.",
            "Настройка базы данных участников, QR и этапов.",
            "Реализация защиты от повторной активации QR.",
            "Реализация админ-команд и выгрузок.",
            "Тестирование пользовательского и админского сценариев.",
            "Развертывание бота на сервере, настройка и запуск.",
        ],
    )

    h2(doc, "8. Что нужно от заказчика")
    bullets(
        doc,
        [
            "Доступ администратора к сообществу VK.",
            "Токен сообщества VK API и group_id.",
            "Подтверждение периода акции (даты 3 этапов).",
            "Согласование финальных текстов сообщений бота.",
        ],
    )

    h2(doc, "9. Сроки реализации")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Этап"
    hdr[1].text = "Содержание"
    hdr[2].text = "Срок"

    rows = [
        ("Подготовка", "Финализация сценария и структуры данных", "1 день"),
        ("Разработка", "Логика бота, QR, этапы, админ-функции", "2-3 дня"),
        ("Тестирование", "Проверка пользовательских и админских сценариев", "1 день"),
        ("Запуск", "Выгрузка на сервер и боевая проверка", "1 день"),
    ]
    for stage, content, term in rows:
        cells = table.add_row().cells
        cells[0].text = stage
        cells[1].text = content
        cells[2].text = term

    doc.add_paragraph("Итоговый срок: 5-6 рабочих дней.")

    h2(doc, "10. Стоимость")
    doc.add_paragraph(
        "Стоимость реализации под ключ: 15 000 руб. "
        "В стоимость входит разработка, настройка, запуск на сервере и базовое сопровождение после запуска."
    )
    doc.add_paragraph(
        "Ежемесячная стоимость сервера, как правило, существенно ниже стоимости конструктора ботов."
    )

    h2(doc, "11. Готовые тексты бота")
    doc.add_paragraph("Приветствие:")
    doc.add_paragraph(
        "Добро пожаловать в розыгрыш Avatar Arena Omsk VR.\n"
        "Сканируйте QR-коды и получайте шансы на приз.\n"
        "Нажмите кнопку ниже, чтобы начать."
    )
    doc.add_paragraph("Кнопка: Принять участие")
    doc.add_paragraph("")
    doc.add_paragraph("Подтверждение активации:")
    doc.add_paragraph("QR успешно активирован. Вам начислен 1 шанс.\nВаше текущее количество шансов: {N}.")
    doc.add_paragraph("")
    doc.add_paragraph("Если QR уже использован:")
    doc.add_paragraph("Этот QR-код уже был активирован ранее.")
    doc.add_paragraph("")
    doc.add_paragraph("Если QR не найден:")
    doc.add_paragraph("QR-код не найден. Проверьте код и повторите попытку.")
    doc.add_paragraph("")
    doc.add_paragraph("Напоминание о финале этапа:")
    doc.add_paragraph("Скоро подведение итогов этапа №{stage}. Ваши шансы уже участвуют в розыгрыше.")
    doc.add_paragraph("")
    doc.add_paragraph("Сообщение победителю:")
    doc.add_paragraph("Поздравляем. Вы стали победителем этапа №{stage}. С вами свяжется менеджер клуба.")

    h2(doc, "12. Итог для заказчика")
    bullets(
        doc,
        [
            "Полностью рабочий бот с прозрачной механикой розыгрыша.",
            "Контроль всех этапов через админ-функции.",
            "Автоматический учет шансов по QR.",
            "Запуск без сторонних сервисов розыгрыша.",
        ],
    )

    out = "Final_TZ_AvatarArena_QR_3Stages.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
