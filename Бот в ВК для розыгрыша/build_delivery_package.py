from __future__ import annotations

import csv
import secrets
import shutil
from pathlib import Path

import qrcode
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


GROUP_DOMAIN = "avatararenaomskvr"
TOKENS_COUNT = 200


def build_link(token: str) -> str:
    return f"https://vk.me/{GROUP_DOMAIN}?ref={token}"


def generate_qr_image(link: str, output_path: Path) -> None:
    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=3,
    )
    qr.add_data(link)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    img.save(output_path)


def create_readme(path: Path) -> None:
    content = """Avatar Arena VK Raffle Package

Что внутри:
- app/main.py: backend бота (FastAPI + SQLite + VK callback)
- app/requirements.txt: зависимости
- app/.env.example: пример переменных окружения
- qr/codes_for_bot.csv: список токенов и ссылок
- qr/png/: PNG-файлы QR-кодов для печати/выдачи
- qr/qr_print_sheets.docx: листы для печати QR (8 на страницу)

Быстрый запуск:
1) cd app
2) pip install -r requirements.txt
3) set переменные окружения из .env.example
4) python main.py
5) открыть http://127.0.0.1:8000/docs
"""
    path.write_text(content, encoding="utf-8")


def main() -> None:
    project_dir = Path(__file__).resolve().parent
    root_dir = project_dir.parent
    package_dir = root_dir / "AvatarArena_VK_Raffle_Package"
    app_dir = package_dir / "app"
    qr_dir = package_dir / "qr"
    qr_png_dir = qr_dir / "png"

    if package_dir.exists():
        shutil.rmtree(package_dir)

    app_dir.mkdir(parents=True, exist_ok=True)
    qr_png_dir.mkdir(parents=True, exist_ok=True)

    shutil.copy2(project_dir / "main.py", app_dir / "main.py")
    (app_dir / "requirements.txt").write_text("fastapi\nuvicorn\nqrcode[pil]\npython-docx\n", encoding="utf-8")
    (app_dir / ".env.example").write_text(
        "\n".join(
            [
                "VK_GROUP_TOKEN=your_group_token_here",
                "VK_CONFIRMATION_TOKEN=your_confirmation_token_here",
                "VK_ADMIN_IDS=123456789",
                "PD_POLICY_URL=https://example.com/personal-data-policy",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    rows: list[dict[str, str]] = []
    for i in range(1, TOKENS_COUNT + 1):
        token = f"QR-{secrets.token_urlsafe(8)}"
        link = build_link(token)
        file_name = f"{i:04d}_{token}.png"
        file_path = qr_png_dir / file_name
        generate_qr_image(link, file_path)
        rows.append({"index": str(i), "token": token, "link": link, "png_file": f"png/{file_name}"})

    with (qr_dir / "codes_for_bot.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["index", "token", "link", "png_file"])
        writer.writeheader()
        writer.writerows(rows)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run("Avatar Arena Omsk VR - QR листы для печати")
    t.bold = True
    t.font.size = Pt(14)
    doc.add_paragraph("")

    per_page = 8
    cols = 2
    rows_per_page = 4
    for start in range(0, len(rows), per_page):
        chunk = rows[start : start + per_page]
        table = doc.add_table(rows=rows_per_page, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for tr in table.rows:
            for cell in tr.cells:
                cell.width = Cm(9)

        for idx, row in enumerate(chunk):
            r = idx // cols
            c = idx % cols
            cell = table.cell(r, c)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = p.add_run(f"Купон #{row['index']}")
            rr.bold = True
            p2 = cell.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.add_paragraph()
            cell.add_paragraph(f"Токен: {row['token']}")
            cell.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p = cell.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_run = img_p.add_run()
            img_run.add_picture(str(qr_dir / row["png_file"]), width=Cm(4.2))

        if start + per_page < len(rows):
            doc.add_page_break()

    doc.save(qr_dir / "qr_print_sheets.docx")
    create_readme(package_dir / "README.txt")
    print(package_dir)


if __name__ == "__main__":
    main()
