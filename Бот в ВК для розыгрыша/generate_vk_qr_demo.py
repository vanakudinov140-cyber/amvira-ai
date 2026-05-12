import csv
from pathlib import Path

import qrcode
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


GROUP_DOMAIN = "avatararenaomskvr"
OUTPUT_DIR = Path("vk_qr_demo_pack")
QR_DIR = OUTPUT_DIR / "qr_png"
COUNT = 20


def build_link(token: str) -> str:
    # vk.me/<domain> opens chat with community, ref carries token.
    return f"https://vk.me/{GROUP_DOMAIN}?ref={token}"


def generate_qr_image(link: str, path: Path) -> None:
    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(link)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    img.save(path)


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    QR_DIR.mkdir(exist_ok=True)

    rows = []
    for i in range(1, COUNT + 1):
        token = f"VRDEMO{i:04d}"
        link = build_link(token)
        file_name = f"{token}.png"
        file_path = QR_DIR / file_name
        generate_qr_image(link, file_path)
        rows.append(
            {
                "token": token,
                "link": link,
                "qr_file": f"qr_png/{file_name}",
            }
        )

    csv_path = OUTPUT_DIR / "vk_qr_demo_map.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["token", "link", "qr_file"])
        writer.writeheader()
        writer.writerows(rows)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("DEMO: QR-коды для запуска в VK")
    r.bold = True
    r.font.size = Pt(16)

    subtitle = doc.add_paragraph(
        "Скан QR открывает диалог с сообществом и передает уникальный ref-токен."
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    for row in rows[:8]:
        doc.add_paragraph(f"Токен: {row['token']}")
        doc.add_paragraph(row["link"])
        doc.add_picture(str(OUTPUT_DIR / row["qr_file"]), width=Cm(4.0))
        doc.add_paragraph("")

    doc.add_paragraph("Полный список QR: папка qr_png. Карта соответствий: vk_qr_demo_map.csv.")
    doc.save(OUTPUT_DIR / "vk_qr_demo_preview.docx")

    print(str(OUTPUT_DIR.resolve()))


if __name__ == "__main__":
    main()
