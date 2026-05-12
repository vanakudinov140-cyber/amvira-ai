import csv
import math
import os
import zipfile
from pathlib import Path

import qrcode
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
import sqlite3


GROUP_ID = os.getenv("TARGET_GROUP_ID", "157843327").strip()
BASE = Path(".").resolve()
SRC_DB = BASE / "qr_6000_pack" / "raffle_6000_for_amvera.db"
OUT = BASE / f"qr_6000_pack_group{GROUP_ID}"
IMG_DIR = OUT / "images"
PRINT_DIR = OUT / "print_docs"


def main() -> None:
    OUT.mkdir(exist_ok=True)
    IMG_DIR.mkdir(exist_ok=True)
    PRINT_DIR.mkdir(exist_ok=True)

    conn = sqlite3.connect(SRC_DB)
    cur = conn.cursor()
    tokens = [r[0] for r in cur.execute("SELECT token FROM qr_tokens ORDER BY id").fetchall()]
    conn.close()

    if len(tokens) != 6000:
        raise RuntimeError(f"Expected 6000 tokens, got {len(tokens)}")

    # CSV with production links
    csv_path = OUT / f"qr_links_6000_group{GROUP_ID}.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["index", "token", "link", "png_file"])
        for i, t in enumerate(tokens, 1):
            link = f"https://vk.me/club{GROUP_ID}?start={t}&ref={t}"
            w.writerow([i, t, link, f"{t}.png"])

    # PNG regeneration with new group link
    for t in tokens:
        link = f"https://vk.me/club{GROUP_ID}?start={t}&ref={t}"
        qrcode.make(link).save(IMG_DIR / f"{t}.png")

    # ZIP images
    images_zip = OUT / f"qr_images_6000_group{GROUP_ID}.zip"
    with zipfile.ZipFile(images_zip, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(IMG_DIR.glob("*.png")):
            z.write(p, arcname=p.name)

    # A4 print docs: 12 per page
    page_w = 21
    page_h = 29.7
    margin = 0.7
    usable_w = page_w - 2 * margin
    cell_w = Cm(usable_w / 3)
    cell_h = Cm(6.6)
    per_doc = 1200

    for part in range(5):
        sub = tokens[part * per_doc : (part + 1) * per_doc]
        d = Document()
        s = d.sections[0]
        s.page_width = Cm(page_w)
        s.page_height = Cm(page_h)
        s.left_margin = Cm(margin)
        s.right_margin = Cm(margin)
        s.top_margin = Cm(margin)
        s.bottom_margin = Cm(margin)

        pages = math.ceil(len(sub) / 12)
        for pg in range(pages):
            page_tokens = sub[pg * 12 : (pg + 1) * 12]
            table = d.add_table(rows=4, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for row in table.rows:
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                row.height = cell_h
            for i, tk in enumerate(page_tokens):
                r = i // 3
                c = i % 3
                cell = table.cell(r, c)
                cell.width = cell_w
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p0 = cell.paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = p0.add_run("КУПОН УЧАСТНИКА РОЗЫГРЫША\n")
                rr.bold = True
                rr.font.size = Pt(7.5)
                run = p0.add_run()
                run.add_picture(str(IMG_DIR / f"{tk}.png"), width=Cm(4.25))
                p2 = cell.add_paragraph(tk)
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p2.runs[0].font.size = Pt(6.5)
            if pg < pages - 1:
                d.add_page_break()

        d.save(PRINT_DIR / f"QR_PRINT_A4_12_group{GROUP_ID}_part{part + 1}.docx")

    # SQL import (tokens unchanged)
    sql_path = OUT / f"qr_tokens_import_6000_group{GROUP_ID}.sql"
    with sql_path.open("w", encoding="utf-8") as f:
        f.write(f"-- Import 6000 QR tokens (links target club{GROUP_ID} in csv/png)\n")
        f.write("INSERT OR IGNORE INTO qr_tokens (token, issued_for_receipt, issued_at, active_for_next_stage)\n")
        f.write("VALUES\n")
        vals = [f"('{t}', 'RCP-{i:04d}', datetime('now'), 1)" for i, t in enumerate(tokens, 1)]
        f.write(",\n".join(vals))
        f.write(";\n")

    # DB copy for Amvera
    db_out = OUT / f"raffle_6000_for_amvera_group{GROUP_ID}.db"
    with open(SRC_DB, "rb") as src, open(db_out, "wb") as dst:
        dst.write(src.read())

    # Full package zip
    pack_zip = BASE / f"qr_6000_pack_group{GROUP_ID}.zip"
    with zipfile.ZipFile(pack_zip, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, arcname=str(p.relative_to(BASE)))

    print(f"created={OUT}")
    print(f"tokens={len(tokens)}")
    print(f"zip={pack_zip}")


if __name__ == "__main__":
    main()

