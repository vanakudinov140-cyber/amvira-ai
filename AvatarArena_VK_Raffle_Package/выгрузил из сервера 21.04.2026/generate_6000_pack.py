import csv
import math
import secrets
import sqlite3
import string
import zipfile
from pathlib import Path

import qrcode
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


COUNT = 6000
GROUP_ID = "236814060"
BASE = Path(".").resolve()
OUT = BASE / "qr_6000_pack"
IMG_DIR = OUT / "images"
PRINT_DIR = OUT / "print_docs"


def gen_tokens(n: int) -> list[str]:
    chars = string.ascii_uppercase + string.digits
    seen = set()
    tokens: list[str] = []
    while len(tokens) < n:
        token = "QR-" + "".join(secrets.choice(chars) for _ in range(11))
        if token in seen:
            continue
        seen.add(token)
        tokens.append(token)
    return tokens


def write_csv(tokens: list[str]) -> Path:
    csv_path = OUT / "qr_links_6000.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["index", "token", "link", "png_file"])
        for i, t in enumerate(tokens, 1):
            link = f"https://vk.me/club{GROUP_ID}?start={t}&ref={t}"
            w.writerow([i, t, link, f"{t}.png"])
    return csv_path


def write_sql(tokens: list[str]) -> Path:
    sql_path = OUT / "qr_tokens_import_6000.sql"
    with sql_path.open("w", encoding="utf-8") as f:
        f.write("-- Import 6000 QR tokens\n")
        f.write("INSERT OR IGNORE INTO qr_tokens (token, issued_for_receipt, issued_at, active_for_next_stage)\n")
        f.write("VALUES\n")
        vals = []
        for i, t in enumerate(tokens, 1):
            rec = f"RCP-{i:04d}"
            vals.append(f"('{t}', '{rec}', datetime(\"now\"), 1)")
        f.write(",\n".join(vals))
        f.write(";\n")
    return sql_path


def generate_png(tokens: list[str]) -> None:
    for t in tokens:
        link = f"https://vk.me/club{GROUP_ID}?start={t}&ref={t}"
        qrcode.make(link).save(IMG_DIR / f"{t}.png")


def zip_images() -> Path:
    zip_path = OUT / "qr_images_6000.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(IMG_DIR.glob("*.png")):
            z.write(p, arcname=p.name)
    return zip_path


def make_print_docs(tokens: list[str]) -> None:
    page_w = 21
    page_h = 29.7
    margin = 0.7
    usable_w = page_w - 2 * margin
    cell_w = Cm(usable_w / 3)
    cell_h = Cm(6.6)
    per_doc = 1200  # 100 pages * 12

    for part in range(5):
        start = part * per_doc
        end = min((part + 1) * per_doc, len(tokens))
        sub = tokens[start:end]

        d = Document()
        s = d.sections[0]
        s.page_width = Cm(page_w)
        s.page_height = Cm(page_h)
        s.left_margin = Cm(margin)
        s.right_margin = Cm(margin)
        s.top_margin = Cm(margin)
        s.bottom_margin = Cm(margin)

        pages = math.ceil(len(sub) / 12)
        for pg in range(pages):
            page_tokens = sub[pg * 12 : (pg + 1) * 12]
            t = d.add_table(rows=4, cols=3)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for row in t.rows:
                row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                row.height = cell_h
            for i, tk in enumerate(page_tokens):
                r = i // 3
                c = i % 3
                cell = t.cell(r, c)
                cell.width = cell_w
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p0 = cell.paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = p0.add_run("КУПОН УЧАСТНИКА РОЗЫГРЫША\n")
                rr.bold = True
                rr.font.size = Pt(7.5)
                run = p0.add_run()
                run.add_picture(str(IMG_DIR / f"{tk}.png"), width=Cm(4.25))
                p2 = cell.add_paragraph(tk)
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p2.runs[0].font.size = Pt(6.5)
            if pg < pages - 1:
                d.add_page_break()

        d.save(PRINT_DIR / f"QR_PRINT_A4_12_part{part + 1}.docx")


def build_amvera_db(tokens: list[str]) -> Path:
    db_template = BASE / "raffle.db"
    db_out = OUT / "raffle_6000_for_amvera.db"

    src_conn = sqlite3.connect(db_template)
    dst_conn = sqlite3.connect(db_out)
    src_conn.backup(dst_conn)
    src_conn.close()

    cur = dst_conn.cursor()
    cur.execute("DELETE FROM draw_winners")
    cur.execute("DELETE FROM draws")
    cur.execute("DELETE FROM activations")
    cur.execute("DELETE FROM pending_tokens")
    cur.execute("UPDATE stages SET status='planned', starts_at=NULL, ends_at=NULL")
    cur.execute("DELETE FROM qr_tokens")
    cur.executemany(
        "INSERT INTO qr_tokens (token, issued_for_receipt, issued_at, active_for_next_stage) VALUES (?, ?, datetime('now'), 1)",
        [(t, f"RCP-{i:04d}") for i, t in enumerate(tokens, 1)],
    )
    dst_conn.commit()
    dst_conn.close()
    return db_out


def zip_full_pack() -> Path:
    pack_zip = BASE / "qr_6000_pack.zip"
    with zipfile.ZipFile(pack_zip, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, arcname=str(p.relative_to(BASE)))
    return pack_zip


def main() -> None:
    OUT.mkdir(exist_ok=True)
    IMG_DIR.mkdir(exist_ok=True)
    PRINT_DIR.mkdir(exist_ok=True)

    tokens = gen_tokens(COUNT)
    write_csv(tokens)
    write_sql(tokens)
    generate_png(tokens)
    zip_images()
    make_print_docs(tokens)
    db_out = build_amvera_db(tokens)
    zip_path = zip_full_pack()

    print(f"created_pack={OUT}")
    print(f"images={len(tokens)}")
    print("print_docs=5")
    print(f"amvera_db={db_out.name}")
    print(f"zip={zip_path}")


if __name__ == "__main__":
    main()

