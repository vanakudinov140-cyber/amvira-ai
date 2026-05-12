from datetime import datetime
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches


BASE_DIR = Path(__file__).resolve().parent
IMAGES_DIR = BASE_DIR / "images"
OUTPUT_FILE = BASE_DIR / "ТЗ_бот_UGC_PRIME.docx"
EMOJI_FONT = "Segoe UI Emoji"


def apply_emoji_font(doc: Document) -> None:
    normal_style = doc.styles["Normal"]
    normal_style.font.name = EMOJI_FONT
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), EMOJI_FONT)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = EMOJI_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), EMOJI_FONT)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_lines(doc: Document, lines: list[str]) -> None:
    for line in lines:
        doc.add_paragraph(line)


def add_html_paragraph(doc: Document, text: str) -> None:
    if text == "":
        doc.add_paragraph("")
        return

    paragraph = doc.add_paragraph()
    bold = False
    underline = False
    italic = False

    tokens = re.split(r"(<\/?[bui]>)", text)
    for token in tokens:
        if token == "<b>":
            bold = True
            continue
        if token == "</b>":
            bold = False
            continue
        if token == "<u>":
            underline = True
            continue
        if token == "</u>":
            underline = False
            continue
        if token == "<i>":
            italic = True
            continue
        if token == "</i>":
            italic = False
            continue
        if token:
            run = paragraph.add_run(token)
            run.bold = bold
            run.underline = underline
            run.italic = italic


def add_html_lines(doc: Document, lines: list[str]) -> None:
    for line in lines:
        add_html_paragraph(doc, line)


def add_image_if_exists(doc: Document, image_path: Path, caption: str) -> None:
    doc.add_paragraph(caption)
    if image_path.exists():
        doc.add_picture(str(image_path), width=Inches(4.8))
    else:
        doc.add_paragraph(f"[Файл не найден: {image_path}]")


def main() -> None:
    doc = Document()

    add_heading(doc, "UGC PRIME — финальное ТЗ и описание логики бота", level=0)
    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph("Документ для заказчика: актуальная структура, тексты, кнопки, визуальные материалы и правила работы бота.")

    add_heading(doc, "1. Общая логика бота", level=1)
    add_lines(
        doc,
        [
            "- Бот работает на aiogram + SQLite.",
            "- Часовой пояс планировщика: Europe/Moscow (МСК).",
            "- Ежедневные напоминания: 16:00 и 20:00 по МСК.",
            "- Учет публикаций ведется по ссылкам, отправленным пользователями.",
            "- Главные сущности БД: users, submissions.",
        ],
    )

    add_heading(doc, "2. Кнопки и навигация (актуальная схема)", level=1)
    add_lines(
        doc,
        [
            "Стартовые кнопки:",
            "• Узнать, чем полезен БОТ ДИСЦИПЛИНЫ📔",
            "• 🏰как повышается мой уровень?",
            "",
            "Кнопка второго шага (после объяснения уровня):",
            "• все поняла! готова быть в прайме⭐️⭐️⭐️",
            "",
            "Основное меню:",
            "• 📶Твой прогресс",
            "• 🔱Твой рейтинг",
            "• 🏰Я опубликовала контент",
            "• 🥺Нет ресурса на контент",
            "",
            "Админ-меню:",
            "• 📊 Статистика",
            "• 📥 Excel",
        ],
    )

    add_heading(doc, "3. Тексты сообщений (актуальные)", level=1)
    add_heading(doc, "3.1 /start", level=2)
    add_html_lines(
        doc,
        [
            "Привет!🏰🩰",
            "Это бот дисциплины <b>UGC PRIME</b>!",
            "",
            "<i>Он будет помогать тебе держать регулярность и доводить до результата через систему!🗝️</i>",
            "",
            "<b>ТВОЯ ЗАДАЧА:</b> <u>после публикации поста/reels отправлять ссылку на опубликованный контент в этот бот!</u>",
            "",
            "Важно: аккаунт должен быть <b>открытым!</b>",
            "📔 чтобы бот мог фиксировать твои публикации!",
        ],
    )

    add_heading(doc, "3.2 Описание бота (после кнопки «Узнать…»)", level=2)
    add_html_lines(
        doc,
        [
            "⭐️<b>БОТ ДИСЦИПЛИНЫ — ТВОЙ ДРУГ НА 2 МЕСЯЦА!</b>",
            "<u>Теперь он автоматически:</u>",
            "🗝️отслеживает твои reels и посты",
            "🗝️фиксирует твой прогресс",
            "🗝️подсказывает, на каком уровне регулярности ты находишься сейчас:",
            "1. <b>ЛЕНТЯЙКА💌</b>",
            "2. <b>В ИГРЕ🦸🏼‍♀️</b>",
            "3. <b>ТЫ В ПРАЙМЕ🏰🕯️</b>",
            "🗝️считает количество опубликованного контента в день и за неделю — подводит итоги недели",
            "🗝️отправляет напоминание о том, что сейчас необходимо опубликовать reels/пост",
            "<b><u>Обязательно:</u></b> поставь уведомление на этот бот — это твой личный трекер⭐️⭐️⭐️",
        ],
    )

    add_heading(doc, "3.3 Сообщение после кнопки «🏰как повышается мой уровень?»", level=2)
    add_html_lines(
        doc,
        [
            "Для повышения твоего уровня в игре, тебе необходимо <b><u>РЕГУЛЯРНО</u></b> публиковать контент (reels и посты) и отправлять это в <b><u>БОТ ДИСЦИПЛИНЫ</u></b>🏰",
            "Во вкладке 📶<u>Твой прогресс</u> ты будешь видеть сколько единиц контента ты опубликовала и собирать 🗝️",
            "🌟<b>Дней подряд</b>: показывает сколько дней подряд ты публикуешь контент",
            "Для того чтобы быть в ПРАЙМЕ🌟 необходимо публиковать 2 единицы контента в день (reels, посты) и собирать ключики",
            "Помни - регулярность. это важно и не забывай присылать ссылки на свои reels и посты!!",
        ],
    )

    add_heading(doc, "3.4 Кнопка под этим сообщением", level=2)
    add_lines(
        doc,
        [
            "все поняла! готова быть в прайме⭐️⭐️⭐️",
        ],
    )

    add_heading(doc, "3.5 Сообщение после «все поняла! готова быть в прайме⭐️⭐️⭐️»", level=2)
    add_html_lines(
        doc,
        [
            "<b><u>Ты в игре!</u></b> Нажимай кнопку ниже после публикации 👇🏼",
        ],
    )

    add_heading(doc, "3.6 Ответ при отправке ссылки", level=2)
    add_html_lines(
        doc,
        [
            "Отлично, присылай ссылку⭐️",
            "После успешной фиксации ссылки: <u>Ты умничка!!🤍</u> (с картинкой success).",
        ],
    )

    add_heading(doc, "3.7 Текст «Нет ресурса на контент»", level=2)
    add_html_lines(
        doc,
        [
            "<b>Аделина, я рядом 🤍</b> (имя динамическое)",
            "И правда понимаю тебя",
            "Но иногда мы просто прячемся в «потом»",
            "<b>Давай сегодня бережно к себе:</b>",
            "— совсем немного",
            "— без усложнений",
            "— без идеальной картинки",
            "Этого уже хватит",
            "<u>Ты уже молодец ✨</u>",
        ],
    )

    add_heading(doc, "4. Логика уровней и прогресса", level=1)
    add_html_lines(
        doc,
        [
            "Уровни по количеству ссылок за текущий день:",
            "• 0 ссылок -> ЛЕНТЯЙКА 💌",
            "• 1 ссылка -> В ИГРЕ 🦸🏼‍♀️",
            "• 2 и более -> ТЫ В ПРАЙМЕ 🏰🕯️",
            "",
            "Блок «📶Твой прогресс» (фактический формат):",
            "<b>Сегодня:</b> X/2",
            "🗝️:  ⚪🔴🔴 (индикатор из 3 кружков)",
            "🗝️Ключи: <b>N</b>",
            "🏰<b>Уровень:</b> ...",
            "📓<b>Всего публикаций за 7 дней:</b> ...",
            "🌟<b>Дней подряд:</b> ...",
            "Ты уже в игре, но <b><u>не выпадай</u></b>🤍",
            "",
            "<b>Индикатор кружков:</b> старт 🔴🔴🔴, при публикациях превращаются в ⚪.",
        ],
    )

    add_heading(doc, "5. Рейтинг (Топ-5)", level=1)
    add_lines(
        doc,
        [
            "Формат вывода:",
            "🔱Топ-5 ПРАЙМОВЫХ результатов за всё время!",
            "1 место: Имя — число",
            "2 место: ...",
            "3 место: ...",
            "4 место: ...",
            "5 место: ...",
            "Поздравляем учениц!🗝️⭐️",
            "Продолжай — ты тоже можешь быть здесь!",
            "",
            "Имена подставляются из users.first_name.",
        ],
    )

    add_heading(doc, "6. Напоминания (по МСК)", level=1)
    add_lines(
        doc,
        [
            "16:00 — текст: время выложить контент!!🌟 + картинка reminder_16.jpg",
            "20:00 — текст: время выложить контент!!🌟 + картинка reminder_20.jpg",
            "Напоминания отправляются участницам, у которых за текущий день count == 0.",
        ],
    )

    add_heading(doc, "7. Экспорт Excel для админа", level=1)
    add_lines(
        doc,
        [
            "Выгружаемые поля:",
            "• User ID",
            "• Имя",
            "• Сегодня (шт)",
            "• За 7 дней (шт)",
            "• Всего публикаций",
            "• Дни подряд",
            "• Лучший день",
            "• Челлендж (день)",
            "• Последняя активность",
            "• Уровень сегодня",
        ],
    )

    add_heading(doc, "8. Приложение: используемые изображения", level=1)
    add_image_if_exists(doc, IMAGES_DIR / "welcome.jpg", "A) welcome.jpg (стартовое сообщение)")
    add_image_if_exists(doc, IMAGES_DIR / "success.jpg", "B) success.jpg (после успешной ссылки)")
    add_image_if_exists(doc, IMAGES_DIR / "no_energy.jpg", "C) no_energy.jpg (кнопка «Нет ресурса на контент»)")
    add_image_if_exists(doc, IMAGES_DIR / "reminder_16.jpg", "D) reminder_16.jpg (напоминание 16:00)")
    add_image_if_exists(doc, IMAGES_DIR / "reminder_20.jpg", "E) reminder_20.jpg (напоминание 20:00)")

    add_heading(doc, "9. Приложение: эмодзи-словарь интерфейса", level=1)
    add_lines(
        doc,
        [
            "🏰 — блок дисциплины/контента",
            "📶 — кнопка и блок прогресса",
            "🔱 — рейтинг",
            "🗝️ — ключи/система",
            "🌟 — стрик/мотивация",
            "🥺 — «Нет ресурса»",
            "🔴/⚪ — индикатор публикаций дня",
            "💌 / 🦸🏼‍♀️ / 🕯️ — уровни и стилистика сообщений",
        ],
    )

    apply_emoji_font(doc)
    doc.save(str(OUTPUT_FILE))
    print(f"Saved: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
